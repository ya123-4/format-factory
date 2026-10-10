"""Generate docx/xlsx test fixtures using the bundled python (python-docx/openpyxl)."""
import sys

from docx import Document
import openpyxl

OUT = r"C:\Users\ybc12\Documents\deepseek-harness\default-workspace\format-factory\tests\fixtures"

d = Document()
d.add_paragraph("格式工厂集成测试")
d.add_paragraph("第二行中文内容")
d.add_paragraph("third line of text")
d.save(OUT + r"\hello.docx")
print("hello.docx written")

wb = openpyxl.Workbook()
ws = wb.active
ws["A1"] = "名称"
ws["B1"] = "数量"
ws["A2"] = "苹果"
ws["B2"] = 42
wb.save(OUT + r"\data.xlsx")
print("data.xlsx written")
