"""End-to-end integration tests (Task 16) — real ffmpeg + LibreOffice Kit."""
import os
import subprocess
import sys
import threading
import time
from pathlib import Path

import pytest

from app.config import Config, DEFAULTS
from app.engine.converter import Converter, required_tools
from app.engine.document import DocConverter
from app.engine.image import ImageConverter
from app.engine.media import MediaEngine
from app.engine.models import ConversionJob, InputKind, JobStatus
from app.engine.runner import Cancelled, atomic_output

FIXTURES = Path(__file__).parent / "fixtures"


def _media_readable(ffmpeg, path) -> bool:
    r = subprocess.run(
        [ffmpeg, "-v", "error", "-i", str(path), "-f", "null", "-"],
        capture_output=True, text=True,
    )
    return r.returncode == 0


def _convert_one(ffmpeg, src, out_dir, fmt):
    job = ConversionJob(src_path=str(src), out_dir=str(out_dir), target_fmt=fmt)
    media = MediaEngine(ffmpeg, "", _cfg(out_dir), None)
    conv = Converter(
        jobs=[job],
        engines={InputKind.VIDEO: media.convert, InputKind.AUDIO: media.convert},
        prober=None,
        config=_cfg(out_dir),
    )
    return conv.run(), job


def _cfg(out_dir):
    from app.config import Config

    return Config(out_dir=str(out_dir), conflict="overwrite")


def test_video_mp4_to_mkv_copy(make_sample_media, ffmpeg, tmp_path):
    src = make_sample_media("video", tmp_path / "in.mp4")
    summary, job = _convert_one(ffmpeg, src, tmp_path, "mkv")
    assert summary.done == 1
    assert job.status == JobStatus.DONE
    assert Path(job.final_path).exists()
    assert _media_readable(ffmpeg, job.final_path)


def test_video_mp4_to_webm_reencode(make_sample_media, ffmpeg, tmp_path):
    src = make_sample_media("video", tmp_path / "in.mp4")
    summary, job = _convert_one(ffmpeg, src, tmp_path, "webm")
    assert summary.done == 1
    assert _media_readable(ffmpeg, job.final_path)


def test_audio_flac_to_mp3(make_sample_media, ffmpeg, tmp_path):
    src = make_sample_media("audio", tmp_path / "in.flac")
    summary, job = _convert_one(ffmpeg, src, tmp_path, "mp3")
    assert summary.done == 1
    assert _media_readable(ffmpeg, job.final_path)


def test_image_bmp_to_png_and_jpg(make_sample_media, ffmpeg, tmp_path):
    src = make_sample_media("image", tmp_path / "in.bmp")
    jobs = [
        ConversionJob(src_path=str(src), out_dir=str(tmp_path), target_fmt=fmt)
        for fmt in ("png", "jpg")
    ]
    eng = ImageConverter(ffmpeg, None, "high")
    conv = Converter(
        jobs=jobs,
        engines={InputKind.IMAGE: eng.convert},
        prober=None,
        config=_cfg(tmp_path),
    )
    summary = conv.run()
    assert summary.done == 2
    for fmt in ("png", "jpg"):
        out = tmp_path / f"in.{fmt}"
        assert out.exists() and out.stat().st_size > 0


def test_document_docx_to_pdf(tmp_path, kit_works):
    if not kit_works:
        pytest.skip("LibreOffice Kit backend blocked in DSH sandbox (spawn EPERM)")
    src = FIXTURES / "hello.docx"
    job = ConversionJob(src_path=str(src), out_dir=str(tmp_path), target_fmt="pdf")
    doc = DocConverter(DEFAULTS["kit_node"], DEFAULTS["kit_cli"], None)
    conv = Converter(
        jobs=[job],
        engines={InputKind.DOCUMENT: doc.convert},
        prober=None,
        config=_cfg(tmp_path),
    )
    summary = conv.run()
    assert summary.done == 1
    out = Path(job.final_path)
    assert out.exists() and out.stat().st_size > 0


def test_document_xlsx_to_csv(tmp_path, kit_works):
    if not kit_works:
        pytest.skip("LibreOffice Kit backend blocked in DSH sandbox (spawn EPERM)")
    src = FIXTURES / "data.xlsx"
    job = ConversionJob(src_path=str(src), out_dir=str(tmp_path), target_fmt="csv")
    doc = DocConverter(DEFAULTS["kit_node"], DEFAULTS["kit_cli"], None)
    conv = Converter(
        jobs=[job],
        engines={InputKind.DOCUMENT: doc.convert},
        prober=None,
        config=_cfg(tmp_path),
    )
    summary = conv.run()
    assert summary.done == 1
    text = Path(job.final_path).read_text(encoding="utf-8", errors="replace")
    assert "苹果" in text


def test_vdat_mp4_variant_end_to_end(make_sample_media, ffmpeg, tmp_path):
    mp4 = make_sample_media("video", tmp_path / "real.mp4")
    vdat = tmp_path / "rec01.vdat"
    vdat.write_bytes(mp4.read_bytes())
    summary, job = _convert_one(ffmpeg, vdat, tmp_path, "mp4")
    assert summary.done == 1
    assert _media_readable(ffmpeg, job.final_path)


def test_vdat_hls_playlist_end_to_end(make_sample_media, ffmpeg, tmp_path):
    # 模拟 cosblaster 布局：.vdat 是 M3U8 清单，分片在 <name>.vdat_contents/
    src = make_sample_media("video", tmp_path / "src.mp4")
    segdir = tmp_path / "play.vdat_contents"
    segdir.mkdir()
    r = subprocess.run(
        [ffmpeg, "-y", "-hide_banner", "-loglevel", "error",
         "-i", str(src), "-c", "copy",
         "-hls_time", "1", "-hls_list_size", "0",
         "-hls_segment_filename", str(segdir / "seg%d.ts"),
         str(tmp_path / "play.m3u8")],
        capture_output=True, text=True,
    )
    assert r.returncode == 0, r.stderr
    # ffmpeg 生成的清单段引用是裸文件名；改写为 cosblaster 风格（相对清单的子目录引用）
    m3u8 = tmp_path / "play.m3u8"
    m3u8.write_text(
        m3u8.read_text(encoding="utf-8").replace(
            "seg0.ts", "play.vdat_contents/seg0.ts"
        ),
        encoding="utf-8",
    )
    vdat = tmp_path / "play.vdat"
    m3u8.replace(vdat)
    summary, job = _convert_one(ffmpeg, vdat, tmp_path, "mp4")
    assert summary.done == 1
    assert _media_readable(ffmpeg, job.final_path)


def test_vdat_hls_missing_segments_fails_cleanly(make_sample_media, ffmpeg, tmp_path):
    # 只有清单没有分片：应 FAILED（InputCorrupt/ConversionFailed），不崩溃
    vdat = tmp_path / "play.vdat"
    vdat.write_text(
        "#EXTM3U\n#EXT-X-VERSION:3\n#EXTINF:1.0,\nplay.vdat_contents/seg0.ts\n",
        encoding="utf-8",
    )
    summary, job = _convert_one(ffmpeg, vdat, tmp_path, "mp4")
    assert summary.failed == 1
    assert job.status == JobStatus.FAILED
    assert job.error


def test_folder_recursive_keeps_rel_dirs(make_sample_media, ffmpeg, tmp_path):
    inner = tmp_path / "视频" / "旅行"
    src = make_sample_media("video", inner / "rec.mp4")
    job = ConversionJob(
        src_path=str(src),
        out_dir=str(tmp_path / "输出"),
        target_fmt="mkv",
        rel_subdir="视频/旅行",
    )
    media = MediaEngine(ffmpeg, "", _cfg(tmp_path / "输出"), None)
    conv = Converter(
        jobs=[job],
        engines={InputKind.VIDEO: media.convert},
        prober=None,
        config=Config(out_dir=str(tmp_path / "输出"), conflict="overwrite"),
    )
    summary = conv.run()
    assert summary.done == 1
    assert (tmp_path / "输出" / "视频" / "旅行" / "rec.mkv").exists()


def test_skip_policy_integration(make_sample_media, ffmpeg, tmp_path):
    src = make_sample_media("video", tmp_path / "in.mp4")
    (tmp_path / "in.mkv").write_bytes(b"exists")
    job = ConversionJob(src_path=str(src), out_dir=str(tmp_path), target_fmt="mkv")
    media = MediaEngine(ffmpeg, "", _cfg(tmp_path), None)
    conv = Converter(
        jobs=[job],
        engines={InputKind.VIDEO: media.convert},
        prober=None,
        config=Config(out_dir=str(tmp_path), conflict="skip"),
    )
    summary = conv.run()
    assert summary.skipped == 1
    assert job.status == JobStatus.SKIPPED


def test_stop_interrupts_and_no_part_left(tmp_path):
    dst = tmp_path / "out.bin"
    stop = threading.Event()
    # cmd 契约：最后一个参数是输出文件（atomic_output 会替换为 .part 临时名）
    cmd = [sys.executable, "-c", "import time; time.sleep(60)", str(dst)]
    errors = []

    def worker():
        try:
            atomic_output(cmd, str(dst), stop)
        except Cancelled:
            errors.append("cancelled")

    t = threading.Thread(target=worker)
    t.start()
    time.sleep(1.0)
    stop.set()
    t.join(timeout=40)
    assert errors == ["cancelled"]
    assert not list(tmp_path.glob("*.part.*"))
    assert not dst.exists()


def test_required_tools_doc_only():
    job = ConversionJob(src_path="a.docx", out_dir="out", target_fmt="pdf")
    assert required_tools([job]) == {"soffice"}
    # PDF 源走进程内引擎，不要求 soffice
    job2 = ConversionJob(src_path="a.pdf", out_dir="out", target_fmt="docx")
    assert required_tools([job2]) == set()


def _make_pdf(tmp_path):
    import fitz  # PyMuPDF

    out = tmp_path / "in.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "FormatFactory PDF test line")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "second page content")
    doc.save(str(out))
    doc.close()
    return out


def _convert_doc(src, tmp_path, fmt):
    job = ConversionJob(src_path=str(src), out_dir=str(tmp_path), target_fmt=fmt)
    doc = DocConverter("", "", None)  # PDF 路径不需要 Kit
    conv = Converter(
        jobs=[job],
        engines={InputKind.DOCUMENT: doc.convert},
        prober=None,
        config=Config(out_dir=str(tmp_path), conflict="overwrite"),
    )
    return conv.run(), job


def test_document_pdf_to_docx(tmp_path):
    src = _make_pdf(tmp_path)
    summary, job = _convert_doc(src, tmp_path, "docx")
    assert summary.done == 1
    out = Path(job.final_path)
    assert out.exists() and out.stat().st_size > 1000
    import zipfile

    assert zipfile.is_zipfile(out)
    with zipfile.ZipFile(out) as z:
        document_xml = z.read("word/document.xml").decode("utf-8")
    assert "FormatFactory" in document_xml
    assert "second page" in document_xml


def test_document_pdf_to_txt(tmp_path):
    src = _make_pdf(tmp_path)
    summary, job = _convert_doc(src, tmp_path, "txt")
    assert summary.done == 1
    text = Path(job.final_path).read_text(encoding="utf-8")
    assert "FormatFactory" in text
    assert "second page" in text


def _make_table_pdf(tmp_path):
    import fitz

    out = tmp_path / "table.pdf"
    doc = fitz.open()
    page = doc.new_page(width=400, height=300)
    # 2x2 表格：外框 + 内线 + 单元格文本
    for x in (80, 180, 280):
        page.draw_line((x, 80), (x, 200), color=(0, 0, 0), width=1)
    for y in (80, 140, 200):
        page.draw_line((80, y), (280, y), color=(0, 0, 0), width=1)
    page.insert_text((100, 120), "Name")
    page.insert_text((200, 120), "Value")
    page.insert_text((100, 180), "Apple")
    page.insert_text((200, 180), "42")
    doc.save(str(out))
    doc.close()
    return out


def test_document_pdf_to_xlsx(tmp_path):
    src = _make_table_pdf(tmp_path)
    summary, job = _convert_doc(src, tmp_path, "xlsx")
    assert summary.done == 1
    out = Path(job.final_path)
    assert out.exists() and out.stat().st_size > 500
    from openpyxl import load_workbook

    wb = load_workbook(out)
    sheets = wb.sheetnames
    assert len(sheets) >= 1
    all_text = " ".join(
        str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value
    )
    assert "Apple" in all_text


def test_document_pdf_to_pptx(tmp_path):
    src = _make_pdf(tmp_path)
    summary, job = _convert_doc(src, tmp_path, "pptx")
    assert summary.done == 1
    out = Path(job.final_path)
    assert out.exists() and out.stat().st_size > 500
    from pptx import Presentation

    prs = Presentation(str(out))
    assert len(prs.slides) == 2  # PDF 两页 → 两张幻灯片
    for slide in prs.slides:
        pics = [s for s in slide.shapes if s.shape_type == 13]  # PICTURE
        assert len(pics) == 1


def test_document_scanned_pdf_ocr(tmp_path):
    # 构造扫描件：文本 PDF 渲染成图片，再把图片包成无文本层的 PDF
    src = _make_pdf(tmp_path)
    import fitz

    text_doc = fitz.open(str(src))
    pix = text_doc[0].get_pixmap(dpi=150)
    text_doc.close()
    scanned = tmp_path / "scanned.pdf"
    out_doc = fitz.open()
    img_page = out_doc.new_page(width=pix.width, height=pix.height)
    img_page.insert_image(img_page.rect, pixmap=pix)
    out_doc.save(str(scanned))
    out_doc.close()
    summary, job = _convert_doc(scanned, tmp_path, "txt")
    assert summary.done == 1
    text = Path(job.final_path).read_text(encoding="utf-8")
    assert "FormatFactory" in text  # OCR 识别出渲染的文字
